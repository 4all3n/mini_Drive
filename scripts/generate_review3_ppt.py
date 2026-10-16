#!/usr/bin/env python3
"""
Generate Review 3 PPT for Intelligent Cloud Drive.
Replicates the exact format of Review 2 PPT (Intelligent_Cloud_Drive_Review2.pptx).

Format spec (extracted from Review 2):
  - Slide size: 20 x 11.25 inches (widescreen)
  - Layout: "Blank" for all slides
  - Header: Dark rectangle (full width, ~2.82 inches tall) at top, centered white bold 48pt Times New Roman
  - Body: TextBox ~0.5in from left, ~3.5in from top, ~18.5in wide, ~7in tall
  - Body font: 24pt Times New Roman black (regular), bold for emphasis
  - Tables: Full width, starting below header
  - References: 18pt
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ============================================================================
# CONSTANTS (matched to Review 2 PPT)
# ============================================================================
SLIDE_WIDTH = 18288000    # 20 inches
SLIDE_HEIGHT = 10287000   # 11.25 inches

# Header bar
HDR_LEFT = 0
HDR_TOP = 0
HDR_WIDTH = 18287280
HDR_HEIGHT = 2577960
HDR_BG_COLOR = RGBColor(0x1A, 0x23, 0x7E)  # dark navy/indigo
HDR_FONT_SIZE = Pt(48)
HDR_FONT_COLOR = RGBColor(0xF0, 0xF0, 0xF0)
HDR_FONT_NAME = "Times New Roman"

# Body text box
BODY_LEFT = 548640
BODY_TOP = 3200400
BODY_WIDTH = 16915680
BODY_HEIGHT = 6400080
BODY_FONT_SIZE = Pt(24)
BODY_FONT_COLOR = RGBColor(0x00, 0x00, 0x00)
BODY_FONT_NAME = "Times New Roman"

# Table positions
TBL_LEFT = 457200
TBL_TOP = 2743200
TBL_WIDTH = 17373240

# References
REF_FONT_SIZE = Pt(18)

# Paths
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
RESULTS_DIR = os.path.join(PROJECT_ROOT, "model", "results")
OUTPUT_PATH = os.path.join(PROJECT_ROOT, "IDP", "Intelligent_Cloud_Drive_Review3.pptx")
TEMPLATE_PATH = os.path.join(PROJECT_ROOT, "IDP", "Intelligent_Cloud_Drive_Review2.pptx")


def create_prs():
    """Create presentation using Review 2 as template for theme/background."""
    prs = Presentation(TEMPLATE_PATH)
    # Remove all existing slides
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].get(qn('r:id'))
        prs.part.drop_rel(rId)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
    return prs


def get_blank_layout(prs):
    """Find the Blank layout."""
    for layout in prs.slide_layouts:
        if layout.name == "Blank":
            return layout
    return prs.slide_layouts[6]  # fallback


def add_header(slide, title_text):
    """Add the dark header bar with white centered title."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, HDR_LEFT, HDR_TOP, HDR_WIDTH, HDR_HEIGHT
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = HDR_BG_COLOR
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title_text.upper()
    run.font.bold = True
    run.font.size = HDR_FONT_SIZE
    run.font.color.rgb = HDR_FONT_COLOR
    run.font.name = HDR_FONT_NAME

    # Vertically center text
    tf.paragraphs[0].space_before = Pt(30)


def add_body_textbox(slide, left=BODY_LEFT, top=BODY_TOP, width=BODY_WIDTH, height=BODY_HEIGHT):
    """Add a body text box and return the text frame."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, height
    )
    shape.fill.background()
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    return tf


def add_body_para(tf, text, bold=False, font_size=BODY_FONT_SIZE, color=BODY_FONT_COLOR, first=False):
    """Add a paragraph to a text frame."""
    if first:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.color.rgb = color
    run.font.name = BODY_FONT_NAME
    run.font.bold = bold
    p.space_after = Pt(6)
    return p


def add_body_mixed(tf, parts, first=False):
    """Add a paragraph with mixed bold/normal segments.
    parts = [(text, bold), (text, bold), ...]
    """
    if first:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    for text, bold in parts:
        run = p.add_run()
        run.text = text
        run.font.size = BODY_FONT_SIZE
        run.font.color.rgb = BODY_FONT_COLOR
        run.font.name = BODY_FONT_NAME
        run.font.bold = bold
    p.space_after = Pt(6)
    return p


def add_table(slide, rows_data, col_widths=None, top=TBL_TOP, left=TBL_LEFT, width=TBL_WIDTH, total_height=5800000):
    """Add a table to the slide. rows_data is list of lists of strings."""
    n_rows = len(rows_data)
    n_cols = len(rows_data[0])
    row_height = Emu(int(total_height / n_rows))

    tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, row_height * n_rows)
    tbl = tbl_shape.table

    # Style cells
    for ri, row_data in enumerate(rows_data):
        for ci, cell_text in enumerate(row_data):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(cell_text)
            run.font.name = BODY_FONT_NAME
            if ri == 0:
                run.font.bold = True
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                # Header row background
                cell_fill = cell.fill
                cell_fill.solid()
                cell_fill.fore_color.rgb = HDR_BG_COLOR
            else:
                run.font.bold = False
                run.font.size = Pt(14)
                run.font.color.rgb = BODY_FONT_COLOR
                if ri % 2 == 0:
                    cell_fill = cell.fill
                    cell_fill.solid()
                    cell_fill.fore_color.rgb = RGBColor(0xE8, 0xEA, 0xF6)

    return tbl_shape


def add_image_centered(slide, img_path, top=3000000, max_width=12000000, max_height=6500000):
    """Add an image centered horizontally."""
    if not os.path.exists(img_path):
        # Add placeholder text instead
        tf = add_body_textbox(slide, BODY_LEFT, top, BODY_WIDTH, 1000000)
        add_body_para(tf, f"[Chart: {os.path.basename(img_path)} — file not found]", bold=True, first=True)
        return
    from PIL import Image
    img = Image.open(img_path)
    img_w, img_h = img.size
    aspect = img_w / img_h

    # Fit within max bounds
    if max_width / aspect <= max_height:
        w = max_width
        h = int(w / aspect)
    else:
        h = max_height
        w = int(h * aspect)

    left = (SLIDE_WIDTH - w) // 2
    slide.shapes.add_picture(img_path, left, top, w, h)


# ============================================================================
# SLIDE BUILDERS
# ============================================================================
def slide_title(prs, layout):
    """Slide 1: Title slide."""
    slide = prs.slides.add_slide(layout)
    # Big centered title
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = HDR_BG_COLOR
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True

    # Title
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(80)
    run = p.add_run()
    run.text = "INTELLIGENT CLOUD DRIVE"
    run.font.bold = True
    run.font.size = Pt(54)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = HDR_FONT_NAME

    # Subtitle
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = "AI-Powered File Management System\nwith Secure Distributed Architecture"
    run2.font.bold = False
    run2.font.size = Pt(32)
    run2.font.color.rgb = RGBColor(0xC0, 0xC0, 0xC0)
    run2.font.name = HDR_FONT_NAME

    # Review line
    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(40)
    run3 = p3.add_run()
    run3.text = "REVIEW 3 — Implementation Progress & Empirical Results"
    run3.font.bold = True
    run3.font.size = Pt(28)
    run3.font.color.rgb = RGBColor(0xFF, 0xD5, 0x4F)
    run3.font.name = HDR_FONT_NAME

    # Student info
    p4 = tf.add_paragraph()
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(50)
    run4 = p4.add_run()
    run4.text = "Student: Allen  |  Guide: [Guide Name]\nDepartment of MCA  |  3rd Semester  |  September 2026"
    run4.font.bold = False
    run4.font.size = Pt(22)
    run4.font.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
    run4.font.name = HDR_FONT_NAME


def slide_agenda(prs, layout):
    """Slide 2: Agenda."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "AGENDA")
    tf = add_body_textbox(slide)
    items = [
        "1. Quick Recap — Problem Statement & Objectives",
        "2. System Architecture (Final)",
        "3. Technology Stack",
        "4. Module Descriptions with Algorithms",
        "5. Dataset Description",
        "6. AI Model — Empirical Evaluation Results",
        "7. Per-Category Performance Analysis",
        "8. Cross-Validation Stability",
        "9. Security Implementation (JWT + BCrypt)",
        "10. Implementation Progress Summary",
        "11. Key Code Highlights",
        "12. Challenges & Solutions",
        "13. Work Remaining (Phase 4)",
        "14. References",
    ]
    for i, item in enumerate(items):
        add_body_para(tf, item, bold=False, first=(i == 0))


def slide_problem(prs, layout):
    """Slide 3: Problem Statement."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "PROBLEM STATEMENT & OBJECTIVES")
    tf = add_body_textbox(slide)
    add_body_mixed(tf, [("The Problem:", True)], first=True)
    add_body_para(tf, "• Data Explosion: Global data creation expected to reach 181 ZB by 2025 (IDC). Manual file organization is unsustainable.")
    add_body_para(tf, "• Security Gaps: Traditional cloud systems use plain-text passwords, session cookies, and lack per-user file isolation.")
    add_body_para(tf, "• No Intelligence: Existing lightweight cloud drives do not auto-categorize files, forcing users into manual folder creation.")
    add_body_para(tf, "")
    add_body_mixed(tf, [("Our Objectives:", True)])
    add_body_para(tf, "1. Design a multi-tier cloud file storage system with a secure REST API")
    add_body_para(tf, "2. Implement JWT-based stateless authentication with BCrypt password hashing")
    add_body_para(tf, "3. Build an AI microservice that classifies files into 6 categories using NLP")
    add_body_para(tf, "4. Empirically compare Multinomial Naive Bayes vs Linear SVM on ultra-short text")
    add_body_para(tf, "5. Develop a JavaFX desktop client with distributed deployment support")


def slide_architecture(prs, layout):
    """Slide 4: Architecture diagram."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "SYSTEM ARCHITECTURE")
    img_path = os.path.join(PROJECT_ROOT, "IDP", "review2_arch.png")
    add_image_centered(slide, img_path, top=2800000, max_width=16500000, max_height=7000000)


def slide_tech_stack(prs, layout):
    """Slide 5: Tech Stack table."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "TECHNOLOGY STACK")
    rows = [
        ["Layer", "Technology", "Version", "Purpose"],
        ["Frontend", "JavaFX", "13", "Desktop GUI — Login, Register, Dashboard screens"],
        ["Frontend Build", "Maven + javafx-maven-plugin", "0.0.4", "Compile & launch JavaFX modules"],
        ["Backend", "Spring Boot", "4.0.3", "REST API server, dependency injection"],
        ["ORM", "Hibernate JPA", "6.x", "Object-Relational Mapping for entities"],
        ["Database", "MariaDB / MySQL", "10.6+", "User accounts, file metadata storage"],
        ["Security", "Spring Security + JJWT", "6.x / 0.12.6", "JWT auth, BCrypt password hashing"],
        ["AI / ML", "Scikit-Learn", "latest", "TF-IDF, Naive Bayes, Linear SVM"],
        ["AI API", "FastAPI + Uvicorn", "latest", "Python REST API serving ML predictions"],
        ["Serialization", "Joblib", "latest", "Trained model persistence to disk"],
        ["Build Tool", "Maven Wrapper (mvnw)", "3.3.4", "Zero-install reproducible builds"],
    ]
    add_table(slide, rows)


def slide_modules(prs, layout):
    """Slide 6: Modules with Algorithms."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "MODULE DESCRIPTIONS & ALGORITHMS")
    tf = add_body_textbox(slide, top=2900000, height=7200000)
    add_body_mixed(tf, [("Module 1: User Authentication", True)], first=True)
    add_body_para(tf, "• Algorithm: BCrypt (Blowfish, cost factor 10) + JWT (HMAC-SHA256, 24hr expiry)")
    add_body_para(tf, "• Flow: Register → BCrypt hash → DB | Login → BCrypt verify → JWT issued → Bearer header on all requests")
    add_body_para(tf, "")
    add_body_mixed(tf, [("Module 2: File Management", True)])
    add_body_para(tf, "• Upload: Multipart → path traversal sanitization → disk write → DB persist + AI categorization")
    add_body_para(tf, "• Download/Delete: IDOR prevention — JWT username vs uploadedBy field ownership check")
    add_body_para(tf, "")
    add_body_mixed(tf, [("Module 3: AI File Categorization", True)])
    add_body_para(tf, "• Algorithm: TF-IDF Character N-grams (n=2..5) + Linear SVM (CalibratedClassifierCV)")
    add_body_para(tf, "• 6 Categories: Archives, Code, Documents, Executables, Media, System & Logs")
    add_body_para(tf, "• Graceful fallback: if AI offline → file marked 'Uncategorized'")
    add_body_para(tf, "")
    add_body_mixed(tf, [("Module 4: JavaFX Desktop Client", True)])
    add_body_para(tf, "• SceneManager pattern, dynamic ServerConfig.java for distributed deployment")
    add_body_para(tf, "• Screens: LoginScreen, RegisterScreen, DashboardScreen (file grid + AI category)")


def slide_dataset(prs, layout):
    """Slide 7: Dataset Description + chart."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "DATASET DESCRIPTION")

    rows = [
        ["Property", "Value"],
        ["Total Samples", "570"],
        ["Train / Test Split", "456 (80%) / 114 (20%)"],
        ["Categories", "6 (perfectly balanced — 95 samples each)"],
        ["Feature Type", "Filenames (ultra-short text, 5–30 characters)"],
        ["Feature Extraction", "TF-IDF with character n-grams (n = 2 to 5)"],
    ]
    add_table(slide, rows, top=2800000, left=457200, width=8500000)

    # Add chart image on the right side
    img_path = os.path.join(RESULTS_DIR, "dataset_distribution.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, 9500000, 2800000, 8000000, 6000000)


def slide_results_main(prs, layout):
    """Slide 8: Main evaluation results comparison table."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "AI MODEL — EMPIRICAL EVALUATION RESULTS")
    rows = [
        ["Metric", "Multinomial Naive Bayes", "Linear SVM (LinearSVC)", "Winner"],
        ["Test Accuracy", "94.74%", "96.49%", "🏆 SVM (+1.75%)"],
        ["Precision (Macro)", "94.85%", "96.75%", "🏆 SVM (+1.90%)"],
        ["Recall (Macro)", "94.74%", "96.49%", "🏆 SVM (+1.75%)"],
        ["F1-Score (Macro)", "94.71%", "96.48%", "🏆 SVM (+1.77%)"],
        ["5-Fold CV F1 (Mean±SD)", "92.40% (±0.018)", "96.47% (±0.011)", "🏆 SVM (+4.07%)"],
    ]
    add_table(slide, rows, top=2800000, total_height=4200000)

    # Key finding text below table (starts at 7300000, no overlap)
    tf = add_body_textbox(slide, top=7300000, height=2200000)
    add_body_mixed(tf, [
        ("Key Finding: ", True),
        ("Linear SVM consistently outperforms Naive Bayes across all metrics with lower variance in cross-validation (±0.011 vs ±0.018), indicating superior generalization stability.", False)
    ], first=True)


def slide_per_class(prs, layout):
    """Slide 9: Per-category F1 scores + chart."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "PER-CATEGORY F1 SCORE ANALYSIS")

    rows = [
        ["Category", "Support", "NB F1", "SVM F1", "Δ"],
        ["Archives", "19", "0.97", "0.97", "—"],
        ["Code", "19", "0.87", "0.93", "+6%"],
        ["Documents", "19", "0.97", "1.00", "+3%"],
        ["Executables", "19", "0.89", "0.89", "—"],
        ["Media", "19", "0.97", "1.00", "+3%"],
        ["System & Logs", "19", "1.00", "1.00", "—"],
    ]
    add_table(slide, rows, top=2800000, left=457200, width=8000000)

    img_path = os.path.join(RESULTS_DIR, "f1_per_class.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, 9000000, 2800000, 8500000, 6200000)


def slide_confusion_nb(prs, layout):
    """Slide 10: Confusion Matrix — Naive Bayes."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "CONFUSION MATRIX — NAIVE BAYES")
    img_path = os.path.join(RESULTS_DIR, "confusion_matrix_nb.png")
    add_image_centered(slide, img_path, top=2800000, max_width=13000000, max_height=7000000)


def slide_confusion_svm(prs, layout):
    """Slide 11: Confusion Matrix — Linear SVM."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "CONFUSION MATRIX — LINEAR SVM")
    img_path = os.path.join(RESULTS_DIR, "confusion_matrix_svm.png")
    add_image_centered(slide, img_path, top=2800000, max_width=13000000, max_height=7000000)


def slide_cross_validation(prs, layout):
    """Slide 12: Cross-validation stability."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "CROSS-VALIDATION STABILITY (5-FOLD)")

    rows = [
        ["Fold", "Naive Bayes F1", "Linear SVM F1"],
        ["1", "0.948", "0.974"],
        ["2", "0.938", "0.948"],
        ["3", "0.894", "0.954"],
        ["4", "0.919", "0.974"],
        ["5", "0.921", "0.974"],
        ["Mean ± SD", "0.924 ± 0.018", "0.965 ± 0.011"],
    ]
    add_table(slide, rows, top=2800000, left=457200, width=8000000)

    img_path = os.path.join(RESULTS_DIR, "cross_validation_boxplot.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, 9000000, 2800000, 8500000, 6200000)


def slide_security(prs, layout):
    """Slide 13: Security implementation."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "SECURITY IMPLEMENTATION")
    rows = [
        ["Feature", "Implementation", "Purpose"],
        ["Password Storage", "BCrypt (Blowfish, cost 10)", "Adaptive hash — expensive for brute-force"],
        ["Token Auth", "JWT HMAC-SHA256, 24hr expiry", "Stateless — no server sessions or cookies"],
        ["Session Strategy", "STATELESS (Spring Security)", "Immune to CSRF attacks"],
        ["Path Traversal", "cleanPath() + normalize() + prefix check", "Blocks ../../etc/passwd attacks"],
        ["IDOR Prevention", "JWT username vs uploadedBy check", "Users can only access own files"],
        ["File Size Limit", "15 MB max (configurable)", "Prevents DoS via oversized uploads"],
    ]
    add_table(slide, rows, top=2800000)


def slide_jwt_flow(prs, layout):
    """Slide 14: JWT Authentication Flow."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "JWT AUTHENTICATION FLOW")
    img_path = os.path.join(PROJECT_ROOT, "IDP", "jwt_flow_diagram.png")
    add_image_centered(slide, img_path, top=2800000, max_width=16500000, max_height=7000000)


def slide_progress(prs, layout):
    """Slide 15: Implementation progress table."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "IMPLEMENTATION PROGRESS (60%+)")
    rows = [
        ["Module", "Status", "Key Deliverables"],
        ["Repository Cleanup", "✅ Done", "Removed artifacts, updated .gitignore"],
        ["AI Model Pipeline", "✅ Done", "Dataset, training, evaluation, 7 charts, serialized models"],
        ["FastAPI AI Service", "✅ Done", "/predict, /health, /evaluate, /predict/detailed endpoints"],
        ["Backend Security", "✅ Done", "BCrypt + JWT + path traversal + IDOR prevention"],
        ["REST API", "✅ Done", "Register, Login, Upload, List, Download, Delete"],
        ["Frontend Network", "✅ Done", "JWT token storage, Bearer headers, ServerConfig.java"],
        ["Admin Panel", "✅ Done", "Legacy servlet-based user management at /admin"],
        ["Run Scripts", "✅ Done", "Fish, Bash, Windows .bat launch scripts"],
        ["Documentation", "✅ Done", "7 detailed setup & deployment guides"],
        ["Distributed Deploy", "✅ Done", "Windows backend + Linux frontend support"],
    ]
    add_table(slide, rows, top=2700000)


def slide_code_highlights(prs, layout):
    """Slide 16: Key code snippets."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "KEY CODE HIGHLIGHTS")
    tf = add_body_textbox(slide, left=BODY_LEFT, top=2800000, width=BODY_WIDTH, height=7200000)
    tf.margin_left = Inches(0.1)

    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    run0 = p0.add_run()
    run0.text = "1. AI Prediction Endpoint (FastAPI — Python):"
    run0.font.bold = True
    run0.font.size = BODY_FONT_SIZE
    run0.font.name = BODY_FONT_NAME
    run0.font.color.rgb = BODY_FONT_COLOR
    p0.space_after = Pt(4)

    add_body_para(tf, '    @app.post("/predict")', font_size=Pt(18))
    add_body_para(tf, '    def predict_category(data: FileData):', font_size=Pt(18))
    add_body_para(tf, '        prediction = active_model.predict([data.filename])[0]', font_size=Pt(18))
    add_body_para(tf, '        return {"category": prediction}', font_size=Pt(18))
    add_body_para(tf, "")

    add_body_mixed(tf, [("2. JWT Token Generation (Java — Spring Boot):", True)])
    add_body_para(tf, "    return Jwts.builder()", font_size=Pt(18))
    add_body_para(tf, '        .subject(username).issuedAt(now)', font_size=Pt(18))
    add_body_para(tf, '        .expiration(new Date(now + 86400000))  // 24 hours', font_size=Pt(18))
    add_body_para(tf, '        .signWith(secretKey)  // HMAC-SHA256', font_size=Pt(18))
    add_body_para(tf, '        .compact();', font_size=Pt(18))
    add_body_para(tf, "")

    add_body_mixed(tf, [("3. AI-Powered Upload with Path Traversal Prevention:", True)])
    add_body_para(tf, '    String clean = StringUtils.cleanPath(file.getOriginalFilename());', font_size=Pt(18))
    add_body_para(tf, '    Path target = Paths.get(UPLOAD_DIR).resolve(clean).normalize();', font_size=Pt(18))
    add_body_para(tf, '    String category = aiService.analyzeAndCategorize(clean);', font_size=Pt(18))
    add_body_para(tf, '    metadata.setCategory(category); fileRepository.save(metadata);', font_size=Pt(18))


def slide_challenges(prs, layout):
    """Slide 17: Challenges."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "CHALLENGES & SOLUTIONS")
    rows = [
        ["Challenge", "Solution"],
        ["Ultra-short text (5-30 chars)", "Character n-grams (n=2..5) instead of word tokenization"],
        ["Frontend-backend protocol mismatch", "Updated to return JSON with JWT; frontend parses & stores token"],
        ["Path traversal vulnerability", "Triple-layer: cleanPath → normalize → startsWith prefix check"],
        ["IDOR on file download/delete", "Ownership check: JWT username vs uploadedBy field"],
        ["Cross-machine deployment", "ServerConfig.java reads host from system props / env vars"],
        ["AI service offline during upload", "Graceful fallback: categorize as 'Uncategorized' on timeout"],
        ["Run scripts broken via symlinks", "Used realpath(status -f) to resolve canonical path"],
    ]
    add_table(slide, rows, top=2800000)


def slide_remaining(prs, layout):
    """Slide 18: Work remaining."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "WORK REMAINING (PHASE 4 — OCTOBER)")
    rows = [
        ["Task", "Priority", "Estimated Effort"],
        ["UI Polish — Dashboard styling, file icons per category", "Medium", "2–3 days"],
        ["Add file search/filter by category", "Medium", "1–2 days"],
        ["Write 60–70 page final project report", "High", "5–7 days"],
        ["Prepare final demo with live deployment", "High", "1 day"],
        ["Optional: Containerize with Docker", "Low", "1 day"],
        ["Optional: Add file preview (images, PDFs)", "Low", "2 days"],
    ]
    add_table(slide, rows, top=2800000)


def slide_metrics_summary(prs, layout):
    """Slide 19: Metrics summary chart."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "METRICS SUMMARY — VISUAL COMPARISON")
    img_path = os.path.join(RESULTS_DIR, "metrics_summary.png")
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, 800000, 2800000, 10500000, 6800000)

    img_path2 = os.path.join(RESULTS_DIR, "accuracy_comparison.png")
    if os.path.exists(img_path2):
        slide.shapes.add_picture(img_path2, 11800000, 3200000, 5600000, 5600000)


def slide_references(prs, layout):
    """Slide 20: References."""
    slide = prs.slides.add_slide(layout)
    add_header(slide, "REFERENCES")
    tf = add_body_textbox(slide, top=3000000, height=6800000)
    refs = [
        '1. Joachims, T. (1998). "Text categorization with Support Vector Machines." Proc. 10th European Conference on Machine Learning.',
        '2. Kowsari, K. et al. (2019). "Text Classification Algorithms: A Survey." Information, 10(4), 150.',
        '3. Jones, M. et al. (2015). "JSON Web Token (JWT)." RFC 7519, IETF.',
        '4. Provos, N. & Mazières, D. (1999). "A Future-Adaptable Password Scheme." USENIX Annual Technical Conference.',
        '5. Pedregosa, F. et al. (2011). "Scikit-learn: Machine Learning in Python." JMLR, 12, 2825–2830.',
        '6. Das, M. et al. (2023). "A comparative study on TF-IDF feature weighting method." Int. J. Advanced Research in CS.',
        '7. Prasad, S. R. S. S. S. & Sharma, V. K. (2023). "Enhancing data security: JWT and HMAC integration." IJCAIT.',
        '8. Biryukov, A. et al. (2016). "Argon2: New generation of memory-hard functions." IEEE Euro S&P.',
        '9. Ghemawat, S. et al. (2003). "The Google File System." Proc. 19th ACM SOSP.',
        '10. FastAPI Documentation. Ramírez, S. (2024). https://fastapi.tiangolo.com.',
    ]
    for i, ref in enumerate(refs):
        add_body_para(tf, ref, font_size=REF_FONT_SIZE, first=(i == 0))


def slide_thank_you(prs, layout):
    """Slide 21: Thank you."""
    slide = prs.slides.add_slide(layout)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = HDR_BG_COLOR
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(120)
    run = p.add_run()
    run.text = "THANK YOU"
    run.font.bold = True
    run.font.size = Pt(66)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = HDR_FONT_NAME

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(40)
    run2 = p2.add_run()
    run2.text = "Questions?"
    run2.font.bold = False
    run2.font.size = Pt(36)
    run2.font.color.rgb = RGBColor(0xC0, 0xC0, 0xC0)
    run2.font.name = HDR_FONT_NAME

    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(50)
    run3 = p3.add_run()
    run3.text = "Live Demo Available — Full stack running:\nMariaDB + FastAPI AI + Spring Boot + JavaFX"
    run3.font.bold = False
    run3.font.size = Pt(24)
    run3.font.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
    run3.font.name = HDR_FONT_NAME


# ============================================================================
# MAIN
# ============================================================================
def main():
    print("Creating Review 3 PPT using Review 2 template format...")
    prs = create_prs()
    layout = get_blank_layout(prs)

    slide_title(prs, layout)           # 1
    slide_agenda(prs, layout)          # 2
    slide_problem(prs, layout)         # 3
    slide_architecture(prs, layout)    # 4
    slide_tech_stack(prs, layout)      # 5
    slide_modules(prs, layout)         # 6
    slide_dataset(prs, layout)         # 7
    slide_results_main(prs, layout)    # 8
    slide_per_class(prs, layout)       # 9
    slide_confusion_nb(prs, layout)    # 10
    slide_confusion_svm(prs, layout)   # 11
    slide_cross_validation(prs, layout) # 12
    slide_security(prs, layout)        # 13
    slide_jwt_flow(prs, layout)        # 14
    slide_progress(prs, layout)        # 15
    slide_code_highlights(prs, layout) # 16
    slide_challenges(prs, layout)      # 17
    slide_remaining(prs, layout)       # 18
    slide_metrics_summary(prs, layout) # 19
    slide_references(prs, layout)      # 20
    slide_thank_you(prs, layout)       # 21

    prs.save(OUTPUT_PATH)
    print(f"✅ Review 3 PPT saved to: {OUTPUT_PATH}")
    print(f"   Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
